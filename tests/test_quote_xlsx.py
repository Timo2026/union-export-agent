"""tests/test_quote_xlsx.py — G2b: 报价单 XLSX 附件闭环.

镜像 quote_pdf 铁律: content_sha256 锁报价事实 (不锁字节), draft_only,
openpyxl 缺库 → 显式降级不炸链; XLSX 与 PDF 同源同锁 (同 payload)。
"""
from __future__ import annotations

from pathlib import Path

from services.quote_pdf import render_quote_pdf
from services.quote_xlsx import attach_quote_xlsx, render_quote_xlsx

QUOTE = {
    "unit_price": 128.5,
    "final_price": 12850.0,
    "total_price": 12850.0,
    "currency": "CNY",
    "lead_time_days": 12,
    "_source": "timo-engine-v12",
}

RFQ = {"material": "6061", "quantity": 100, "surface": "as-machined",
       "tolerance_grade": "IT7"}
CUSTOMER = {"name": "Acme", "email": "buyer@acme.cn"}


def _sha_len_ok(s: str) -> bool:
    return isinstance(s, str) and len(s) == 64 and all(c in "0123456789abcdef" for c in s)


# ---- 1) 渲染 ----

def test_render_quote_xlsx_ok(tmp_path):
    r = render_quote_xlsx(QUOTE, context_id="RFQ-G2B-1", out_dir=tmp_path,
                          customer=CUSTOMER, rfq=RFQ)
    assert r["ok"] is True, r
    p = Path(r["path"])
    assert p.exists() and p.stat().st_size > 0
    assert "RFQ-G2B-1" in p.name and p.suffix == ".xlsx"
    assert _sha_len_ok(r["content_sha256"])
    # XLSX = zip 容器 (PK 魔数)
    assert p.read_bytes()[:2] == b"PK"


def test_content_sha_matches_pdf(tmp_path):
    """XLSX 与 PDF 锁同一份报价事实 (同 payload → 同 content_sha256)."""
    x = render_quote_xlsx(QUOTE, context_id="RFQ-G2B-2", out_dir=tmp_path / "x",
                          customer=CUSTOMER, rfq=RFQ)
    p = render_quote_pdf(QUOTE, context_id="RFQ-G2B-2", out_dir=tmp_path / "p",
                         customer=CUSTOMER, rfq=RFQ)
    assert x["ok"] is True, x
    if p["ok"]:  # reportlab 不在CI环境时跳过 PDF 对照
        assert x["content_sha256"] == p["content_sha256"]


def test_render_deterministic_content_sha(tmp_path):
    r1 = render_quote_xlsx(QUOTE, context_id="RFQ-G2B-3", out_dir=tmp_path / "a")
    r2 = render_quote_xlsx(QUOTE, context_id="RFQ-G2B-3", out_dir=tmp_path / "b")
    assert r1["ok"] and r2["ok"]
    assert r1["content_sha256"] == r2["content_sha256"]


def test_render_no_price_returns_noop(tmp_path):
    r = render_quote_xlsx({"currency": "CNY"}, context_id="RFQ-G2B-4", out_dir=tmp_path)
    assert r["ok"] is False and r["reason"] == "no_price"
    assert not list(tmp_path.rglob("*.xlsx"))


def test_render_missing_openpyxl_graceful(tmp_path, monkeypatch):
    import services.quote_xlsx as qx
    monkeypatch.setattr(qx, "_openpyxl_available", lambda: False)
    r = render_quote_xlsx(QUOTE, context_id="RFQ-G2B-5", out_dir=tmp_path)
    assert r["ok"] is False and r["reason"] == "openpyxl_not_installed"
    assert not list(tmp_path.rglob("*.xlsx"))


def test_xlsx_cells_carry_quote_facts(tmp_path):
    """表内数字 = 确定性引擎产出原样 (不重算)."""
    from openpyxl import load_workbook
    r = render_quote_xlsx(QUOTE, context_id="RFQ-G2B-6", out_dir=tmp_path,
                          customer=CUSTOMER, rfq=RFQ)
    assert r["ok"] is True, r
    wb = load_workbook(r["path"])
    ws = wb.active
    rows = [[c.value for c in row] for row in ws.iter_rows()]
    flat = [str(v) for row in rows for v in row if v is not None]
    assert any("Acme" in v for v in flat)
    assert any("128.5" in v for v in flat)      # unit_price 原样
    assert any("12850" in v for v in flat)      # final_price 原样
    assert any("DRAFT" in v.upper() for v in flat)  # 水印/草稿声明在表内


# ---- 2) 草稿附件闭环 ----

def test_attach_quote_xlsx_adds_attachment_to_draft(tmp_path):
    from services.reply import build_reply
    draft = build_reply({"context_id": "RFQ-G2B-7", "rfq": RFQ, "quote": QUOTE,
                         "customer": CUSTOMER, "verification_status": "PASS"},
                        {"status": "PASS"})
    out = attach_quote_xlsx(draft, quote=QUOTE, context_id="RFQ-G2B-7",
                            customer=CUSTOMER, rfq=RFQ, out_dir=tmp_path)
    atts = out["attachments"]
    assert len(atts) == 1
    a = atts[0]
    assert a["kind"] == "quote_xlsx"
    assert a["draft_only"] is True
    assert Path(a["path"]).exists()
    assert _sha_len_ok(a["content_sha256"])
    assert draft["auto_send"] is False and draft["mode"] == "draft_only"
    assert out.get("attachment_error") is None


def test_attach_missing_openpyxl_marks_error_on_draft(tmp_path, monkeypatch):
    import services.quote_xlsx as qx
    monkeypatch.setattr(qx, "_openpyxl_available", lambda: False)
    from services.reply import build_reply
    draft = build_reply({"context_id": "RFQ-G2B-8", "rfq": RFQ, "quote": QUOTE,
                         "customer": CUSTOMER, "verification_status": "PASS"},
                        {"status": "PASS"})
    out = attach_quote_xlsx(draft, quote=QUOTE, context_id="RFQ-G2B-8",
                            customer=CUSTOMER, rfq=RFQ, out_dir=tmp_path)
    assert out["attachments"] == []
    assert out["attachment_error"] == "openpyxl_not_installed"
    assert draft["auto_send"] is False and draft["mode"] == "draft_only"
